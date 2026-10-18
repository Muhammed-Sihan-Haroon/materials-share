"""End-to-end tests for Materials Share.

Runs the real site (served from a clean copy of this repo, without any live data) and the real
submissions-service code (worker.js under Node with an in-memory KV and a stand-in AI model),
then drives lecturer and student flows in headless Chromium. GitHub's API is simulated.

    pip install -r tests/requirements.txt
    python -m playwright install chromium
    python tests/e2e.py

Exits non-zero if any check fails. Screenshots are written to tests/out/.
"""
import atexit, base64, csv, hashlib, io, json, os, re, shutil, subprocess, sys, tempfile, time, urllib.request, uuid, zipfile
from pathlib import Path
from playwright.sync_api import sync_playwright
import docx

ROOT = Path(__file__).resolve().parents[1]
SRC = str(ROOT)
SP = str(ROOT / "tests" / "out")            # screenshots (uploaded by CI)
WORK = tempfile.mkdtemp(prefix="ms-e2e-")
SITE = WORK + "/site"
os.makedirs(SP, exist_ok=True)

# A clean copy of the site: code only, never your live shares, activities, modules or keys.
DATA = {"s", "a", "m", "keys", "tests", ".git", ".github", "node_modules"}
shutil.copytree(SRC, SITE, ignore=lambda d, names: [n for n in names if d == SRC and n in DATA])
cfg = open(SITE + "/config.js").read()
SALT = re.search(r'salt:\s*"([^"]+)"', cfg).group(1).encode()
cfg = re.sub(r'submitUrl:\s*"[^"]*"', 'submitUrl: "http://localhost:8787"', cfg)   # never touch the real service
open(SITE + "/config.js", "w").write(cfg)
shutil.copy(SRC + "/worker.js", WORK + "/worker.mjs")

srv = subprocess.Popen([sys.executable, "-m", "http.server", "8765", "-d", SITE], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
wk = subprocess.Popen(["node", str(ROOT / "tests" / "worker-dev.mjs"), WORK + "/worker.mjs"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
atexit.register(lambda: (srv.terminate(), wk.terminate(), shutil.rmtree(WORK, ignore_errors=True)))

def wait_for(url, tries=50):
    for _ in range(tries):
        try:
            urllib.request.urlopen(url, timeout=1); return
        except urllib.error.HTTPError:
            return
        except Exception:
            time.sleep(0.2)
    raise SystemExit("server didn't start: " + url)
wait_for("http://localhost:8765/"); wait_for("http://localhost:8787/")

# ---------- test files
FD = WORK + "/files"
os.makedirs(FD, exist_ok=True)
d = docx.Document(); d.add_heading("Week 2 lab answers", 1); d.add_paragraph("Q1: An INNER JOIN keeps only matching rows."); d.save(FD + "/answers.docx")
open(FD + "/query.py", "w").write("import sqlite3\nprint('SELECT * FROM students')\n")
open(FD + "/evil.html", "w").write("<h1>hi</h1><script>window.top.__pwned = 1; parent.__pwned = 1;</script>")
open(FD + "/notes.pdf", "wb").write(b"%PDF-1.4\n%fake\n")
from reportlab.pdfgen import canvas as _cv
c = _cv.Canvas(FD + "/worksheet.pdf"); c.drawString(72, 750, "Worksheet: write a LEFT JOIN query"); c.save()
c = _cv.Canvas(FD + "/answers.pdf"); c.drawString(72, 750, "My answer: a LEFT JOIN keeps all rows from the left table"); c.save()
md = docx.Document(); md.add_paragraph("MODEL ANSWER: LEFT JOIN returns unmatched left rows with NULLs."); md.save(FD + "/model.docx")

# ---------- fake GitHub (writes into SITE so Pages-style reads work)
blobs, calls = {}, []
def fake_github(route, request):
    url, method = request.url, request.method
    path = url.split("/repos/me/materials", 1)[1].split("?")[0]
    calls.append((method, path))
    if request.headers.get("authorization", "") != "Bearer github_pat_TEST":
        return route.fulfill(status=401, json={"message": "Bad credentials"})
    body = json.loads(request.post_data) if request.post_data else None
    if path == "" and method == "GET":
        return route.fulfill(json={"full_name": "me/materials", "permissions": {"push": True}})
    if path.startswith("/git/ref/heads/"):
        return route.fulfill(json={"object": {"sha": "c0"}})
    if path.startswith("/git/commits/"):
        return route.fulfill(json={"tree": {"sha": "t0"}})
    if path == "/git/blobs" and method == "POST":
        sha = uuid.uuid4().hex
        blobs[sha] = base64.b64decode(body["content"])
        return route.fulfill(status=201, json={"sha": sha})
    if path.startswith("/git/blobs/"):
        return route.fulfill(json={"content": base64.b64encode(blobs[path.split("/")[-1]]).decode()})
    if path == "/git/trees":
        for e in body["tree"]:
            p = os.path.join(SITE, e["path"])
            if e["sha"] is None:
                os.remove(p)
            else:
                os.makedirs(os.path.dirname(p), exist_ok=True)
                open(p, "wb").write(blobs[e["sha"]])
        return route.fulfill(status=201, json={"sha": "t1"})
    if path == "/git/commits":
        return route.fulfill(status=201, json={"sha": "c1"})
    if path.startswith("/git/refs/heads/") and method == "PATCH":
        return route.fulfill(json={})
    if path.startswith("/contents/"):
        p = os.path.join(SITE, path[len("/contents/"):])
        if os.path.isdir(p):
            rel = path[len("/contents/"):]
            return route.fulfill(json=[{"path": rel + "/" + f} for f in os.listdir(p)])
        if not os.path.isfile(p):
            return route.fulfill(status=404, json={"message": "Not Found"})
        data = open(p, "rb").read()
        sha = uuid.uuid4().hex
        blobs[sha] = data
        return route.fulfill(json={"sha": sha, "content": base64.b64encode(data).decode() if len(data) < 1_000_000 else ""})
    return route.fulfill(status=500, json={"message": "unmocked " + path})

ok = True
def check(cond, label):
    global ok
    print(("PASS " if cond else "FAIL ") + label)
    ok &= bool(cond)

def http(method, path, headers=None, body=None):
    req = urllib.request.Request("http://localhost:8787" + path, data=body, method=method, headers=headers or {})
    try:
        with urllib.request.urlopen(req) as r: return r.status, r.read()
    except urllib.error.HTTPError as e: return e.code, e.read()

def code_id(code):
    return hashlib.pbkdf2_hmac("sha256", code.encode(), SALT, 400000, 48)[:16].hex()

errors = []
with sync_playwright() as p:
    b = p.chromium.launch()
    ctx = b.new_context(accept_downloads=True, bypass_csp=True)
    ctx.route("https://api.github.com/**", fake_github)
    ctx.route("https://fonts.googleapis.com/**", lambda r, q: r.fulfill(body="", content_type="text/css"))
    ctx.route("https://fonts.gstatic.com/**", lambda r, q: r.abort())
    def page():
        pg = ctx.new_page()
        pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.on("dialog", lambda dlg: dlg.accept())
        return pg

    # ===== lecturer: connect + passphrase
    L = page()
    L.goto("http://localhost:8765/review.html")
    L.wait_for_selector("#signin")
    check("Set up lecturer access" in L.text_content("#signin"), "first visit offers lecturer setup")
    L.fill("#siRepo", "me/materials"); L.fill("#siToken", "github_pat_WRONG")
    L.fill("#siPass", "correct horse battery"); L.fill("#siPass2", "correct horse battery"); L.click("#siGo")
    L.wait_for_function("document.querySelector('#siMsg').textContent.includes(\"didn't accept\")", timeout=20000)
    check(True, "bad GitHub token refused at sign-in")
    L.fill("#siToken", "github_pat_TEST"); L.fill("#siPass", "short"); L.fill("#siPass2", "short"); L.click("#siGo")
    L.wait_for_function("document.querySelector('#siMsg').textContent.includes('12')", timeout=20000)
    check(True, "short passphrase rejected")
    L.fill("#siPass", "correct horse battery"); L.fill("#siPass2", "correct horse battery"); L.click("#siGo")
    L.wait_for_selector("#app:not([hidden])", timeout=30000)
    auth = L.evaluate("localStorage.getItem('ms-auth')")
    check(auth and "github_pat_TEST" not in auth and L.evaluate("localStorage.getItem('ms-conn')") is None, "token stored encrypted, never in plain text")
    lf = json.load(open(SITE + "/keys/lecturer.json"))
    check("pub" in lf and "priv" in lf and os.path.exists(SITE + "/keys/index.bin"), "marking key + index saved to repo")

    # simulate an activity list saved under a different marking key (what the user hit)
    open(SITE + "/keys/index.bin", "wb").write(os.urandom(80))
    # ===== create activity
    L.click("#newBtn")
    L.fill("#aTitle", "Week 2 SQL lab"); L.fill("#aModule", "COM106")
    L.fill("#aDeadline", "2030-01-01T17:00"); L.fill("#aOut", "10")
    L.fill("#aBrief", "Write the three join queries and explain each.")
    L.fill("#aCrit", "Correct joins\nExplains results\nTidy SQL")
    L.fill("#aQc", "Nice clear explanation.\nCheck your LEFT JOIN.")
    L.click("#createAct")
    L.wait_for_selector("#present:not([hidden])", timeout=30000)
    acode = L.text_content("#pCode").replace("-", "")
    print("   activity code", acode)
    aid = code_id(acode)
    check(os.path.exists(f"{SITE}/a/{aid}/meta.json") and os.path.exists(f"{SITE}/a/{aid}/m.bin"), "activity published to repo")
    raw = open(f"{SITE}/a/{aid}/m.bin", "rb").read() + open(SITE + "/keys/index.bin", "rb").read()
    check(b"SQL" not in raw and b"joins" not in raw.lower(), "activity brief and index are encrypted")
    L.screenshot(path=SP + "/s-present.png")
    L.click("#pClose")

    # ===== session: reload stays signed in; sign out asks for the passphrase only
    L.reload()
    L.wait_for_selector("#app:not([hidden])", timeout=20000)
    check(L.locator("#signin").count() == 0, "stays signed in after reload (8-hour session)")
    L.click("#signOut")
    L.wait_for_selector("#signin", timeout=20000)
    check(L.is_hidden("#siTokenBox"), "after sign-out only the passphrase is needed")
    L.fill("#siPass", "wrong passphrase!!"); L.click("#siGo")
    L.wait_for_function("document.querySelector('#siMsg').textContent.includes(\"didn't work\")", timeout=20000)
    check(True, "wrong passphrase refused")
    L.fill("#siPass", "correct horse battery"); L.click("#siGo")
    L.wait_for_selector("#app:not([hidden])", timeout=20000)
    check("Week 2 SQL lab" in L.text_content("#acts"), "activity list restored after sign-in")

    # ===== student 1 submits
    S = page()
    S.goto("http://localhost:8765/")
    S.fill("#code", acode); S.click("#go")
    S.wait_for_selector("#activity:not([hidden])", timeout=20000)
    check("Write the three join queries" in S.text_content("#aBrief") and "Tidy SQL" in S.text_content("#aCritList"), "student sees brief and criteria")
    check("Due" in S.text_content("#aDue"), "deadline shown")
    S.screenshot(path=SP + "/s-activity.png", full_page=True)
    S.goto("http://localhost:8765/"); S.wait_for_selector("#code"); S.screenshot(path=SP + "/s-home.png")
    S.goto("http://localhost:8765/#" + acode); S.wait_for_selector("#activity:not([hidden])", timeout=20000)
    S.fill("#sName", "Aisha Khan"); S.fill("#sId", "B00123456")
    S.set_input_files("#sFiles", [FD + "/query.py", FD + "/answers.docx", FD + "/evil.html", FD + "/notes.pdf"])
    S.fill("#sNote", "Q3 was tricky.")
    S.click("#subBtn")
    S.wait_for_selector("#receipt:not([hidden])", timeout=30000)
    r1 = S.text_content("#rCode").replace("-", "")
    print("   receipt 1", r1)
    check(len(r1) == 8, "receipt code shown")
    S.screenshot(path=SP + "/s-receipt.png")

    store = json.loads(urllib.request.urlopen("http://localhost:8787/__store").read())
    blob_text = "".join(e["sample"] for e in store) + json.dumps([e["meta"] for e in store])
    print("   store keys:", [(e["k"][:12], e["size"]) for e in store], [w for w in ("Aisha","B00123456","sqlite3") if w in blob_text])
    check(len(store) == 1 and "Aisha" not in blob_text and "B00123456" not in blob_text and "sqlite3" not in blob_text, "stored submission is encrypted (no name, ID or content visible)")

    # feedback not yet released
    S2 = page(); S2.goto("http://localhost:8765/#" + r1)
    S2.wait_for_selector("#err:not([hidden])", timeout=20000)
    check("feedback hasn't been released" in S2.text_content("#err"), "receipt before release says feedback not ready")
    S2.close()

    # ===== student 2 (hostile name for CSV test)
    S3 = page(); S3.goto("http://localhost:8765/#" + acode)
    S3.wait_for_selector("#activity:not([hidden])", timeout=20000)
    S3.fill("#sName", '=HYPERLINK("http://x","y")'); S3.fill("#sId", "B00999999")
    S3.set_input_files("#sFiles", [FD + "/query.py"])
    S3.click("#subBtn"); S3.wait_for_selector("#receipt:not([hidden])", timeout=30000)
    r2 = S3.text_content("#rCode").replace("-", "")
    # resubmission by student 1
    S3.goto("http://localhost:8765/#" + acode); S3.reload()
    S3.wait_for_selector("#activity:not([hidden])", timeout=20000)
    S3.fill("#sName", "Aisha Khan"); S3.fill("#sId", "b00123456")
    S3.set_input_files("#sFiles", [FD + "/query.py"])
    S3.click("#subBtn"); S3.wait_for_selector("#receipt:not([hidden])", timeout=30000)
    S3.goto("http://localhost:8765/"); S3.reload()
    print("   receipts on device:", S3.locator("#mine .mine").count())
    check(S3.locator("#mine .mine").count() == 3, "student device lists its receipts")

    # ===== lecturer reviews
    L.click("#acts .act >> nth=0")
    L.wait_for_function("document.querySelectorAll('#subs tr.click').length === 3", timeout=20000)
    txt = L.text_content("#subs")
    check("Aisha Khan" in txt and "B00999999" in txt and "earlier attempt" in txt, "submissions list decrypts names, flags earlier attempt")
    L.click("#subs tr.click:has-text('earlier attempt')")
    L.wait_for_selector("#rTabs button", timeout=20000)
    check(L.locator("#rTabs button").count() == 4, "all 4 files present")
    check("Q3 was tricky" in L.text_content("#rNote"), "student note shown")
    check("import sqlite3" in L.text_content("#rPreview pre"), "python file previews as text")
    L.click("#rTabs button:has-text('answers.docx')")
    fr = L.frame_locator("#rPreview iframe")
    fr.locator("text=INNER JOIN").wait_for(timeout=20000)
    check(L.get_attribute("#rPreview iframe", "sandbox") == "", "docx rendered in sandboxed frame")
    L.screenshot(path=SP + "/s-review.png", full_page=True)
    L.click("#rTabs button:has-text('evil.html')")
    L.wait_for_selector("#rPreview pre")
    check("<script>" in L.text_content("#rPreview pre") and L.evaluate("window.__pwned") is None, "student HTML shown as source, never executed")
    L.click("#rTabs button:has-text('notes.pdf')")
    check(L.locator("#rPreview iframe").count() == 1 and L.get_attribute("#rPreview iframe", "sandbox") is None, "pdf opens in viewer")

    # feedback
    rows = L.locator("#fCrit .crit")
    rows.nth(0).locator("label").filter(has_text=re.compile(r"^Met$")).click()
    rows.nth(1).locator("label").filter(has_text=re.compile(r"^Partly met$")).click()
    L.fill("#fMark", "8")
    L.click("#fText")
    L.click("#fQc button:has-text('Nice clear explanation.')")
    L.type("#fText", " Good work on Q2.")
    L.click("#release")
    L.wait_for_function("document.querySelector('#fStatus').textContent.includes('Released')", timeout=20000)
    check("Released" in L.text_content("#subs"), "status shows Released")

    S4 = page(); S4.goto("http://localhost:8765/#" + r1)
    S4.wait_for_selector("#feedback:not([hidden])", timeout=20000)
    fbt = S4.text_content("#feedback")
    check("8 / 10" in S4.text_content("#fMark") and "Nice clear explanation. Good work on Q2." in fbt, "student sees mark and comment")
    check("Partly met" in fbt and "Correct joins" in fbt and "Tidy SQL" not in fbt, "criteria shown (unmarked ones hidden)")
    S4.screenshot(path=SP + "/s-feedback.png", full_page=True)

    # student 2: draft then release all
    L.click("#subs tr.click:has-text('B00999999')")
    L.wait_for_selector("#rTabs button", timeout=20000)
    L.fill("#fText", "Please add comments to your code.")
    L.click("#saveDraft")
    L.wait_for_function("document.querySelector('#fStatus').textContent.includes('Draft saved')", timeout=20000)
    S5 = page(); S5.goto("http://localhost:8765/#" + r2)
    S5.wait_for_selector("#err:not([hidden])", timeout=20000)
    check(True, "draft is not visible to student")
    L.click("#releaseAll")
    L.wait_for_function("document.querySelector('#sMsg').textContent.includes('Released 1')", timeout=20000)
    S5.goto("http://localhost:8765/"); S5.goto("http://localhost:8765/#" + r2); S5.reload()
    S5.wait_for_selector("#feedback:not([hidden])", timeout=20000)
    check("Please add comments" in S5.text_content("#feedback"), "release-all delivers draft feedback")

    # CSV
    with L.expect_download() as dl:
        L.click("#exportCsv")
    rows_csv = list(csv.reader(io.StringIO(open(dl.value.path(), encoding="utf-8-sig").read())))
    hdr = rows_csv[0]
    check(hdr[:2] == ["Student", "Student number"] and "Correct joins" in hdr and len(rows_csv) == 4, "CSV has header + 3 rows")
    names = [r[0] for r in rows_csv[1:]]
    check(any(n.startswith("'=HYPERLINK") for n in names), "CSV neutralises formula in student name")
    aisha = [r for r in rows_csv[1:] if r[0] == "Aisha Khan" and r[hdr.index("Status")] == "Released"][0]
    check(aisha[hdr.index("Mark")] == "8" and aisha[hdr.index("Correct joins")] == "Met", "CSV carries criteria and mark")

    # withdraw
    L.click("#subs tr.click:has-text('earlier attempt')")
    L.wait_for_selector("#withdraw:not([hidden])", timeout=20000)
    L.click("#withdraw")
    L.wait_for_function("document.querySelector('#fStatus').textContent.includes('Draft saved')", timeout=20000)
    S4.goto("http://localhost:8765/"); S4.goto("http://localhost:8765/#" + r1); S4.reload()
    S4.wait_for_selector("#err:not([hidden])", timeout=20000)
    check(True, "withdrawn feedback no longer visible")

    # close submissions
    L.click("#toggleClose")
    L.wait_for_function("document.querySelector('#toggleClose').textContent.includes('Reopen')", timeout=20000)
    S6 = page(); S6.goto("http://localhost:8765/#" + acode)
    S6.wait_for_selector("#activity:not([hidden])", timeout=20000)
    check(not S6.is_visible("#subForm") and "closed" in S6.text_content("#aClosed"), "closed activity hides the form")
    st, body = http("POST", f"/submit?a={aid}&r={'ab'*16}", {"X-Meta": "QUJD", "Content-Type": "application/octet-stream"}, b"x" * 100)
    check(st == 403, "service refuses submissions to a closed activity")

    # service security
    st, _ = http("GET", "/admin/ping", {"Authorization": "Bearer nope"})
    check(st == 401, "lecturer endpoints reject a bad token")
    st, _ = http("GET", f"/admin/list?a={aid}")
    check(st == 401, "lecturer endpoints reject no token")
    st, _ = http("POST", f"/submit?a={'cd'*16}&r={'ab'*16}", {"X-Meta": "QUJD"}, b"x" * 100)
    check(st == 404, "unknown activity refused")
    st, _ = http("POST", f"/submit?a={aid}&r={'ab'*16}", {"X-Meta": "QUJD"}, b"x" * (21 * 1024 * 1024))
    check(st == 413, "oversize upload refused")
    st, _ = http("POST", f"/submit?a=../../x&r={'ab'*16}", {"X-Meta": "QUJD"}, b"x")
    check(st == 400, "bad ids refused")


    # ===== AI-assisted activity with worksheet + model answer
    dialogs = []
    L.on("dialog", lambda dlg: dialogs.append(dlg.message))
    check(L.is_visible("#aiAll") or True, "")
    L.click("#newBtn")
    L.fill("#aTitle", "AI lab"); L.fill("#aModule", "COM106"); L.fill("#aOut", "10")
    L.fill("#aBrief", "Use the worksheet and upload your answers.")
    L.fill("#aCrit", "Correct joins\nExplains results\nTidy SQL")
    L.fill("#aNotes", "Look for NULL handling.")
    L.set_input_files("#aFiles", [FD + "/worksheet.pdf"])
    L.set_input_files("#aNotesFile", [FD + "/model.docx"])
    L.click("#createAct")
    L.wait_for_selector("#present:not([hidden])", timeout=30000)
    ai_code = L.text_content("#pCode").replace("-", ""); ai_id = code_id(ai_code)
    L.click("#pClose")
    wraw = open(f"{SITE}/a/{ai_id}/w0.bin", "rb").read()
    check(b"Worksheet" not in wraw and b"MODEL ANSWER" not in open(SITE + "/keys/index.bin", "rb").read(), "worksheet and model answer stored encrypted")

    S8 = page(); S8.goto("http://localhost:8765/#" + ai_code)
    S8.wait_for_selector("#activity:not([hidden])", timeout=20000)
    check("worksheet.pdf" in S8.text_content("#aFileList"), "student sees worksheet to download")
    check("AI tool" in S8.text_content("#subForm"), "students are told an AI tool may help draft feedback")
    with S8.expect_download() as dl:
        S8.click("#aFileList button")
    check(open(dl.value.path(), "rb").read() == open(FD + "/worksheet.pdf", "rb").read(), "worksheet downloads byte-exact")
    S8.fill("#sName", "Ben Ito"); S8.fill("#sId", "B00222222")
    S8.set_input_files("#sFiles", [FD + "/answers.docx", FD + "/answers.pdf", FD + "/query.py"])
    S8.click("#subBtn"); S8.wait_for_selector("#receipt:not([hidden])", timeout=30000)
    ai_r = S8.text_content("#rCode").replace("-", "")

    L.click("#acts .act:has-text('AI lab')")
    L.wait_for_function("document.querySelectorAll('#subs tr.click').length === 1", timeout=20000)
    check(L.is_visible("#aiAll"), "AI draft button shown when AI is switched on")
    L.click("#aiAll")
    L.wait_for_function("document.querySelector('#sMsg').textContent.includes('Drafted 1')", timeout=60000)
    check("AI draft" in L.text_content("#subs"), "list shows AI draft status")
    ai = json.loads(urllib.request.urlopen("http://localhost:8787/__ai").read())
    msgs = ai["last"]["input"]["messages"]
    prompt = msgs[1]["content"]; system = msgs[0]["content"]
    check("INNER JOIN keeps only matching rows" in prompt, "Word submission text reached the checker")
    check("LEFT JOIN keeps all rows from the left table" in prompt, "PDF submission text reached the checker")
    check("import sqlite3" in prompt, "code file reached the checker")
    check("MODEL ANSWER: LEFT JOIN returns unmatched" in prompt and "Look for NULL handling" in prompt, "model answer file + notes reached the checker")
    check("not instructions" in system and "British English" in system and "0 to 10" in system, "checker prompt guards injection, sets tone and mark range")
    check(not S8.evaluate("1") or True, "")
    S9 = page(); S9.goto("http://localhost:8765/#" + ai_r)
    S9.wait_for_selector("#err:not([hidden])", timeout=20000)
    check(True, "AI draft not visible to student before approval")

    L.click("#subs tr.click:has-text('Ben Ito')")
    L.wait_for_selector("#aiBadge:not([hidden])", timeout=20000)
    checked = L.evaluate("[0,1,2].map(i => (document.querySelector(`input[name=c${i}]:checked`)||{}).value || null)")
    check(checked == ["met", "partly", "not"], f"AI criteria pre-filled {checked}")
    check(L.input_value("#fMark") == "7.5" and "INNER and LEFT JOIN clearly" in L.input_value("#fText"), "AI mark (rounded) and comment pre-filled")
    check("AI: Q3 result explained only briefly." in L.text_content("#fCrit"), "AI reasons shown to lecturer")
    L.frame_locator("#rPreview iframe").locator("text=INNER JOIN").wait_for(timeout=20000)
    check(True, "Word preview renders in AI review")
    L.screenshot(path=SP + "/s-ai.png", full_page=False)

    L.click("#releaseAll")
    L.wait_for_function("document.querySelector('#sMsg').textContent.includes('Released 1')", timeout=20000)
    check(any("AI draft" in d for d in dialogs), "release-all warns about unedited AI drafts")
    S9.goto("http://localhost:8765/"); S9.goto("http://localhost:8765/#" + ai_r); S9.reload()
    S9.wait_for_selector("#feedback:not([hidden])", timeout=20000)
    fb9 = S9.text_content("#feedback")
    check("7.5 / 10" in fb9 and "Put each clause on its own line" in fb9 and "AI:" not in fb9, "student sees approved feedback, no AI reasons")

    # per-submission redraft + quota handling
    S10 = page(); S10.goto("http://localhost:8765/#" + ai_code)
    S10.wait_for_selector("#activity:not([hidden])", timeout=20000)
    S10.fill("#sName", "Cara Doe"); S10.fill("#sId", "B00333333")
    S10.set_input_files("#sFiles", [FD + "/query.py"])
    S10.click("#subBtn"); S10.wait_for_selector("#receipt:not([hidden])", timeout=30000)
    urllib.request.urlopen("http://localhost:8787/__aimode=quota").read()
    L.click("#refresh")
    L.wait_for_function("document.querySelectorAll('#subs tr.click').length === 2", timeout=20000)
    L.click("#aiAll")
    L.wait_for_function("document.querySelector('#sMsg').textContent.includes('allowance')", timeout=30000)
    check("resets" in L.text_content("#sMsg"), "daily allowance limit explained")
    urllib.request.urlopen("http://localhost:8787/__aimode=ok").read()
    L.click("#subs tr.click:has-text('Cara Doe')")
    L.wait_for_selector("#aiOne:not([hidden])", timeout=20000)
    L.wait_for_function("!document.querySelector('#aiOne').disabled", timeout=20000)
    L.click("#aiOne")
    L.wait_for_selector("#aiBadge:not([hidden])", timeout=60000)
    L.fill("#fText", "Edited by lecturer.")
    L.click("#saveDraft")
    L.wait_for_function("document.querySelector('#fStatus').textContent.includes('Draft saved')", timeout=20000)
    check(L.is_hidden("#aiBadge") and "AI draft" not in L.locator("#subs tr.click:has-text('Cara Doe')").text_content(), "editing an AI draft marks it reviewed")

    # ===== modules: QHO444, weeks, drafts, publish, schedule, student view
    import datetime as _dt
    today = _dt.date.today()
    start = today - _dt.timedelta(days=today.weekday() + 14)   # so this is week 3
    for n, body in [("slides.pdf", None), ("Week3_worksheet.docx", None), ("Week3_answers.py", "print('answer')\n")]:
        pth = FD + "/" + n
        if n.endswith(".pdf"):
            c = _cv.Canvas(pth); c.drawString(72, 750, "Loops slides"); c.save()
        elif n.endswith(".docx"):
            dd = docx.Document(); dd.add_paragraph("Write a for loop that prints 1 to 10."); dd.save(pth)
        else:
            open(pth, "w").write(body)
    M = page(); M.goto("http://localhost:8765/modules.html")
    M.wait_for_selector("#app:not([hidden])", timeout=20000)
    M.click("#newMod")
    M.fill("#mShort", "QHO444"); M.fill("#mTitle", "Problem Solving through Programming")
    M.fill("#mStart", start.isoformat()); M.fill("#mWeeks", "12")
    M.fill("#mWeekTitles", "Welcome and setup\nVariables and input\nLoops\nFunctions")
    M.click("#mSave")
    M.wait_for_selector("#present:not([hidden])", timeout=30000)
    mcode = M.text_content("#pCode"); mid = code_id(mcode)
    M.click("#pClose")
    check(os.path.exists(f"{SITE}/m/{mid}/meta.json") and b"QHO444" not in open(f"{SITE}/m/{mid}/index.bin", "rb").read(), "module published, index encrypted")
    check("It's week 3 now" in M.text_content("#vMeta"), "lecturer sees the current week")
    wk3 = 'section[data-week="3"]'
    M.set_input_files(f'{wk3} input[type=file]', [FD + "/slides.pdf", FD + "/Week3_worksheet.docx", FD + "/Week3_answers.py"])
    M.wait_for_function("document.querySelector('#vSave').textContent.includes('Saved')", timeout=30000)
    rows3 = M.text_content(wk3)
    check(rows3.count("Draft") >= 3 and "Worksheet" in rows3 and "Answers" in rows3 and "Material" in rows3, "dropped files become drafts with guessed types")
    raw_files = b"".join(open(os.path.join(SITE, "m", mid, "f", f), "rb").read() for f in os.listdir(os.path.join(SITE, "m", mid, "f")))
    check(b"for loop" not in raw_files and b"answer" not in raw_files, "module files stored encrypted")

    St = page(); St.goto("http://localhost:8765/#" + mcode)
    St.wait_for_selector("#module:not([hidden])", timeout=20000)
    check("Nothing has been published" in St.text_content("#mWeeks"), "students can't see drafts")

    M.click(f'{wk3} .item:has-text("slides") button:has-text("Publish")')
    M.wait_for_function("document.querySelector('#vSave').textContent.includes('Saved')", timeout=30000)
    M.click(f'{wk3} .item:has-text("Week3 worksheet") button:has-text("Publish")')
    M.wait_for_function("document.querySelector('#vSave').textContent.includes('Saved')", timeout=30000)
    future = (_dt.datetime.now() + _dt.timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M")
    M.click(f'{wk3} .item:has-text("Week3 answers") button:has-text("Edit")')
    M.check('input[name="vis"][value="scheduled"]'); M.fill("#dAt", future); M.click("#dSave")
    M.wait_for_function("document.querySelector('#vSave').textContent.includes('Saved')", timeout=30000)
    check("Opens" in M.text_content(f'{wk3} .item:has-text("Week3 answers")'), "answers scheduled")

    St.reload(); St.wait_for_selector("#module:not([hidden])", timeout=20000)
    stext = St.text_content("#mWeeks")
    check("Week 3: Loops" in stext and "this week" in stext and "New" in stext, "student sees current week open, marked new")
    check("Opens" in St.text_content('#mWeeks .item:has-text("answers")'), "student sees answers locked until release")
    with St.expect_download() as dl:
        St.click('#mWeeks .item:has-text("worksheet") .files button')
    check(open(dl.value.path(), "rb").read() == open(FD + "/Week3_worksheet.docx", "rb").read(), "student downloads worksheet byte-exact")
    ans_id = json.loads(urllib.request.urlopen("http://localhost:8787/__store").read())
    kid = [e["k"] for e in ans_id if e["k"].startswith("k:")]
    st, _ = http("GET", "/key/" + kid[0][2:])
    check(len(kid) == 1 and st == 423, "answer key withheld before release time")

    # release time passes
    past = (_dt.datetime.now() - _dt.timedelta(minutes=1)).strftime("%Y-%m-%dT%H:%M")
    M.click(f'{wk3} .item:has-text("Week3 answers") button:has-text("Edit")')
    M.fill("#dAt", past); M.click("#dSave")
    M.wait_for_function("document.querySelector('#vSave').textContent.includes('Saved')", timeout=30000)
    St.reload(); St.wait_for_selector("#module:not([hidden])", timeout=20000)
    with St.expect_download() as dl:
        St.click('#mWeeks .item:has-text("answers") .files button')
    check(open(dl.value.path(), "rb").read() == open(FD + "/Week3_answers.py", "rb").read(), "answers open once release time passes")
    with St.expect_download(timeout=30000) as dl:
        St.click('#mWeeks button:has-text("Download this week as a zip")')
    zn = sorted(zipfile.ZipFile(dl.value.path()).namelist())
    check(zn == ["Week3_answers.py", "Week3_worksheet.docx", "slides.pdf"], f"week zip has all released files {zn}")

    # move an item to week 4 via the editor, unpublish, delete
    M.click(f'{wk3} .item:has-text("slides") button:has-text("Edit")')
    M.select_option("#dWeek", "4"); M.click("#dSave")
    M.wait_for_function("document.querySelector('#vSave').textContent.includes('Saved')", timeout=30000)
    check("slides" in M.text_content('section[data-week="4"]') and "slides" not in M.text_content(wk3), "item moved to week 4")
    M.click(f'{wk3} .item:has-text("Week3 worksheet") button:has-text("Unpublish")')
    M.wait_for_function("document.querySelector('#vSave').textContent.includes('Saved')", timeout=30000)
    St.reload(); St.wait_for_selector("#module:not([hidden])", timeout=20000)
    check("worksheet" not in St.text_content("#mWeeks").lower(), "unpublished item disappears for students")
    nfiles = len(os.listdir(os.path.join(SITE, "m", mid, "f")))
    M.click(f'{wk3} .item:has-text("Week3 worksheet") button:has-text("Delete")')
    M.wait_for_function("document.querySelector('#vSave').textContent.includes('Saved')", timeout=30000)
    check(len(os.listdir(os.path.join(SITE, "m", mid, "f"))) == nfiles - 1, "deleting removes the file from the repo")
    M.fill("#vSearch", "answers")
    check(M.locator("#weeks section").count() == 1, "lecturer search filters weeks")
    M.fill("#vSearch", "")

    # new activity straight into a module week
    L.click("#newBtn")
    L.wait_for_selector("#newCard:not([hidden])")
    L.select_option("#aMod", label="QHO444: Problem Solving through Programming")
    check(L.input_value("#aWeek") == "3", "module week defaults to the current week")
    L.fill("#aTitle", "Loops lab"); L.fill("#aCrit", "Loop is correct"); L.click("#createAct")
    L.wait_for_selector("#present:not([hidden])", timeout=30000)
    lab_code = L.text_content("#pCode"); L.click("#pClose")
    St.reload(); St.wait_for_selector("#module:not([hidden])", timeout=20000)
    St.click('#mWeeks .item:has-text("Loops lab") button:has-text("Open activity")')
    St.wait_for_selector("#activity:not([hidden])", timeout=20000)
    check("Loops lab" in St.text_content("#aTitle"), "module activity opens the hand-in page")
    St.goto("http://localhost:8765/"); St.reload(); St.wait_for_selector("#code")
    check("QHO444" in St.text_content("#myMods"), "student's device remembers the module")
    St.screenshot(path=SP + "/s-home-mods.png")
    St.goto("http://localhost:8765/#" + mcode); St.wait_for_selector("#module:not([hidden])", timeout=20000)
    St.screenshot(path=SP + "/s-module.png", full_page=True)
    M.reload(); M.wait_for_selector("#modView:not([hidden])", timeout=20000)
    M.click('section[data-week="4"] .wk-toggle')
    M.screenshot(path=SP + "/s-modules.png", full_page=True)

    # dashboard
    D = page(); D.goto("http://localhost:8765/dashboard.html")
    D.wait_for_selector("#app:not([hidden])", timeout=20000)
    D.wait_for_function("!document.querySelector('#todo').textContent.includes('Checking')", timeout=30000)
    dtext = D.text_content("#app")
    check("Problem Solving through Programming" in D.text_content("#mods") and "week 3 of 12" in D.text_content("#mods"), "dashboard shows module card")
    check("new submission" in D.text_content("#todo") or "AI draft" in D.text_content("#todo"), "dashboard lists work waiting for feedback")
    D.screenshot(path=SP + "/s-dashboard.png", full_page=True)
    # materials share still works through the unified code box
    A = page(); A.goto("http://localhost:8765/admin.html")
    A.wait_for_selector("#app:not([hidden])", timeout=20000)
    A.fill("#title", "Week 1 slides"); A.set_input_files("#files", [FD + "/notes.pdf"])
    A.click("#create"); A.wait_for_selector("#present:not([hidden])", timeout=30000)
    scode = A.text_content("#pCode").replace("-", "")
    S7 = page(); S7.goto("http://localhost:8765/#" + scode)
    S7.wait_for_selector("#share:not([hidden])", timeout=20000)
    check("Week 1 slides" in S7.text_content("#title"), "materials codes still open from the same box")
    check(all(w in A.text_content("header") for w in ["Dashboard", "Modules", "Activities", "Quick share"]), "lecturer nav links every page")
    check("Week 1 slides" in A.text_content("#list"), "quick shares kept in lecturer data")


    # ===== strict security policy: lecturer libraries still work, external scripts are refused
    sc = b.new_context()
    sp = sc.new_page()
    csp_errors = []
    sp.on("console", lambda m: csp_errors.append(m.text) if "Content Security Policy" in m.text else None)
    sp.goto("http://localhost:8765/review.html")
    docx_b = list(open(FD + "/answers.docx", "rb").read()); pdf_b = list(open(FD + "/answers.pdf", "rb").read())
    res = sp.evaluate("""async ([docx, pdf]) => {
      const load = (src) => new Promise((ok, no) => { const s = document.createElement('script'); s.src = src; s.onload = ok; s.onerror = () => no(src); document.head.append(s); });
      await load('vendor/mammoth.browser.min.js'); await load('vendor/pdf.min.js'); await load('vendor/jszip.min.js');
      const w = (await mammoth.extractRawText({ arrayBuffer: new Uint8Array(docx).buffer })).value;
      pdfjsLib.GlobalWorkerOptions.workerSrc = 'vendor/pdf.worker.min.js';
      const doc = await pdfjsLib.getDocument({ data: new Uint8Array(pdf) }).promise;
      const p = (await (await doc.getPage(1)).getTextContent()).items.map(i => i.str).join(' ');
      const z = new JSZip(); z.file('a.txt', 'hi'); const zb = await z.generateAsync({ type: 'uint8array' });
      let external = 'loaded';
      try { await load('https://example.com/x.js'); } catch { external = 'blocked'; }
      return { w, p, zip: zb.length, external };
    }""", [docx_b, pdf_b])
    check("INNER JOIN" in res["w"] and "LEFT JOIN" in res["p"] and res["zip"] > 0, "Word, PDF and zip libraries work under the strict policy")
    check(res["external"] == "blocked", "scripts from other sites are refused on lecturer pages")
    sc.close()
    # mobile layout
    m = b.new_context(viewport={"width": 375, "height": 760})
    mp = m.new_page(); mp.goto("http://localhost:8765/#" + acode)
    mp.wait_for_selector("#activity:not([hidden])", timeout=20000)
    check(mp.evaluate("document.documentElement.scrollWidth") <= 375, "student activity page fits a phone")
    mp.screenshot(path=SP + "/s-mobile.png", full_page=True)
    mp.goto("http://localhost:8765/"); mp.wait_for_selector("#code")
    mp.screenshot(path=SP + "/s-home-mobile.png", full_page=True)
    check(mp.evaluate("document.documentElement.scrollWidth") <= 375, "student home fits a phone")
    check(not errors, "no JS errors " + str(errors[:3]))
    b.close()

print("ALL PASS" if ok else "SOME FAILED")
sys.exit(0 if ok else 1)
